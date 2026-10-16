"""Document -> provenance-preserving chunks. Parsers run only in the container."""
import re
import subprocess
import tempfile
from pathlib import Path
from .common import digest, normalize, write_jsonl, write_json


def pages(path, ocr=False):
    ext = path.suffix.lower()
    if ext in {'.md', '.markdown', '.txt'}:
        yield 1, path.read_text(encoding='utf-8-sig')
    elif ext == '.docx':
        from docx import Document
        doc = Document(path)
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += '\n' + '\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
        yield 1, text
    elif ext == '.doc':
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(['libreoffice', '-env:UserInstallation=file://' + tmp + '/profile',
                            '--headless', '--convert-to', 'docx', '--outdir', tmp, str(path.resolve())],
                           check=True, capture_output=True, timeout=120)
            yield from pages(Path(tmp) / (path.stem + '.docx'), ocr)
    elif ext == '.pdf':
        from pypdf import PdfReader
        for i, page in enumerate(PdfReader(path).pages):
            text = page.extract_text() or ''
            if len(normalize(text)) < 30 and ocr:
                import pypdfium2 as pdfium
                import pytesseract
                pdf = pdfium.PdfDocument(str(path))
                image = pdf[i].render(scale=2).to_pil()
                text = pytesseract.image_to_string(image, lang='chi_sim+eng')
                image.close(); pdf.close()
            if len(normalize(text)) < 30:
                raise ValueError(f'PDF 第 {i+1} 页无有效文本，请使用 --ocr 或检查源文件')
            yield i+1, text
    else:
        raise ValueError(f'不支持的格式: {ext}')


def ingest(input_dir, output, chunk_size=700, overlap=100, ocr=False):
    if not 0 <= overlap < chunk_size:
        raise ValueError('必须满足 0 <= overlap < chunk_size')
    root = Path(input_dir); rows = []; rejected = []; seen = set()
    for path in sorted(root.rglob('*')):
        if not path.is_file() or path.suffix.lower() not in {'.pdf','.doc','.docx','.md','.markdown','.txt'}:
            continue
        try:
            extracted = list(pages(path, ocr))
            source_id = digest(normalize(' '.join(t for _, t in extracted)))[:24]
            for page, raw in extracted:
                title = next((s.lstrip('# ').strip() for s in raw.splitlines() if s.strip()), path.stem)[:120]
                text = normalize(raw)
                for start in range(0, len(text), chunk_size-overlap):
                    chunk = text[start:start+chunk_size]
                    key = digest(chunk)
                    if len(chunk) < 40 or key in seen:
                        continue
                    seen.add(key)
                    rows.append({'id': key[:24], 'source_id': source_id, 'text': chunk,
                                 'title': title, 'source': str(path.relative_to(root)), 'page': page,
                                 'offset': start, 'content_hash': key})
        except Exception as e:
            rejected.append({'source': str(path.relative_to(root)), 'reason': str(e)})
    write_jsonl(output, rows)
    write_json(str(output) + '.report.json', {'chunks': len(rows), 'rejected': rejected})
    if rejected:
        raise ValueError(f'{len(rejected)} 个文档解析失败，已记录报告；修复后重跑，禁止静默漏文档')
    if not rows:
        raise ValueError('没有可用文本块')
