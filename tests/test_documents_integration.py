"""Real document conversion/OCR integration; runs in the built training image."""
import shutil
import subprocess
from pathlib import Path
import pytest
from aml.documents import pages


@pytest.mark.skipif(not shutil.which('libreoffice') or not shutil.which('tesseract'),reason='Requires container document tools')
def test_word_pdf_legacy_doc_and_scanned_ocr(tmp_path):
    from docx import Document
    import pypdfium2 as pdfium
    doc=Document()
    doc.add_heading('Backup Policy',0)
    doc.add_paragraph('Backup retention is thirty days. A full backup is scheduled every Sunday. Restore requests require approval.')
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Owner';table.cell(0,1).text='Team Blue'
    original=tmp_path/'policy.docx';doc.save(original)
    assert 'Team Blue' in list(pages(original))[0][1]
    for ext in ('pdf','doc'):
        subprocess.run(['libreoffice','-env:UserInstallation=file://'+str(tmp_path/'profile'),
                        '--headless','--convert-to',ext,'--outdir',str(tmp_path),str(original)],check=True,capture_output=True,timeout=60)
        assert 'retention' in ' '.join(t for _,t in pages(tmp_path/f'policy.{ext}')).lower()
    pdf=pdfium.PdfDocument(str(tmp_path/'policy.pdf'))
    image=pdf[0].render(scale=2).to_pil().convert('RGB')
    image.save(tmp_path/'scan.pdf');image.close();pdf.close()
    with pytest.raises(ValueError,match='无有效文本'):list(pages(tmp_path/'scan.pdf',ocr=False))
    assert 'retention' in ' '.join(t for _,t in pages(tmp_path/'scan.pdf',ocr=True)).lower()
